import json

from docx import Document
from docx.shared import Inches
from PIL import Image

from brain_loader.parse import Parser


def test_real_docling_docx_table_figure_and_cache(tmp_path):
    image = tmp_path / "diagram.png"
    Image.new("RGB", (120, 80), color="blue").save(image)
    document = Document()
    document.add_heading("Example document", level=1)
    document.add_paragraph("A document with an intact table and an embedded figure.")
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("Metric", "Value"), ("Alpha", "10"), ("Beta", "20")], strict=True):
        for cell, text in zip(row.cells, values, strict=True):
            cell.text = text
    document.add_picture(str(image), width=Inches(2))
    source = tmp_path / "source.docx"
    document.save(source)
    parser = Parser()
    blocks = parser.parse(source, tmp_path)
    tables = [b for b in blocks if b.kind == "table"]
    figures = [b for b in blocks if b.kind == "figure"]
    assert len(tables) == 1
    assert all(value in tables[0].text for value in ("Metric", "Alpha", "Beta", "10", "20"))
    assert len(figures) == 1
    assert (tmp_path / "figures/0000.png").is_file()
    assert json.loads((tmp_path / "docling.json").read_text(encoding="utf-8"))
    source.unlink()  # cached extraction does not need the source or model inference again
    assert parser.parse(source, tmp_path) == blocks


def test_pdf_checkpoints_preserve_page_numbers_and_resume(tmp_path, monkeypatch):
    from types import SimpleNamespace

    import pypdfium2
    from docling_core.types.doc import DoclingDocument, Size

    class PDF:
        def __init__(self, source):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def __len__(self):
            return 9

    calls = []

    class Converter:
        def convert(self, source, page_range, raises_on_error):
            calls.append(page_range)
            if page_range == (9, 9) and len(calls) == 2:
                raise RuntimeError("interrupted")
            doc = DoclingDocument(name="book")
            for page in range(page_range[0], page_range[1] + 1):
                doc.add_page(page_no=page, size=Size(width=100, height=200))
            return SimpleNamespace(status=SimpleNamespace(value="success"), document=doc)

    monkeypatch.setattr(pypdfium2, "PdfDocument", PDF)
    parser = Parser()
    parser.converter = Converter()
    import pytest

    with pytest.raises(RuntimeError, match="interrupted"):
        parser._pdf(tmp_path / "source.pdf", tmp_path)
    merged = parser._pdf(tmp_path / "source.pdf", tmp_path)
    assert calls == [(1, 8), (9, 9), (9, 9)]
    assert set(merged.pages) == set(range(1, 10))
